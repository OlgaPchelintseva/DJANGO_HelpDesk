from django.contrib import admin
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from django.shortcuts import render
from .models import Client, Ticket, Comment, ImportLog
from .forms import ExcelInputForm
from django.urls import path
from django.utils.html import format_html
from django.utils import timezone
from django.http import HttpResponse
from django.utils.dateparse import parse_datetime

class CommentInline(admin.TabularInline):
    model = Comment
    extra = 0
    readonly_fields = ('created_at', )

@admin.register(Ticket)
class TicketAdmin(admin.ModelAdmin):
    inlines = [CommentInline]

    @admin.display(description='Приоритет')
    def priority_badge(self, obj):
        colors = {
            Ticket.Priority.LOW: 'green',
            Ticket.Priority.NORMAL: 'yellow',
            Ticket.Priority.HIGH: 'orange',
            Ticket.Priority.CRITICAL: 'red',
        }
        labels = {
            Ticket.Priority.LOW: 'Низкий',
            Ticket.Priority.NORMAL: 'Нормальный',
            Ticket.Priority.HIGH: 'Высокий',
            Ticket.Priority.CRITICAL: 'Критичный',
        }
        color = colors.get(obj.priority, 'gray')
        label = labels.get(obj.priority, 'Неизвестно')
        
        return format_html(
            '<span style="background-color: {}; color: white; padding: 2px 6px; border-radius: 3px; font-size: 12px;"><strong>{}</strong></span>',
            color,
            label,
        )

    @admin.display(description='Статус')
    def status_badge(self, obj):
        colors = {
            Ticket.Status.NEW: 'blue',
            Ticket.Status.IN_PROGRESS: 'purple',
            Ticket.Status.WAITING: 'pink',
            Ticket.Status.CLOSED: 'gray',
        }
        labels = {
            Ticket.Status.NEW: 'Новый',
            Ticket.Status.IN_PROGRESS: 'В процессе',
            Ticket.Status.WAITING: 'Ожидание',
            Ticket.Status.CLOSED: 'Закрытый',
        }
        color = colors.get(obj.status, 'black')
        label = labels.get(obj.status, 'Неизвестно')
        return format_html(
            '<span style="background-color: {}; color: white; padding: 2px 6px; border-radius: 3px; font-size: 12px;"><strong>{}</strong></span>',
            color,
            label,
        )
    
    list_display = ('external_id', 'subject', 'client', 'priority_badge', 'status_badge', 'assignee', 'created_at')
    list_filter = ('priority', 'status', 'assignee', 'created_at')
    search_fields = ('subject', 'client__name', 'external_id')
    readonly_fields = ('created_at', 'closed_at')
    change_list_template = 'helpdesk/admin/ticket_change_list.html'

    def save_model(self, request, obj, form, change):
        if obj.status == Ticket.Status.CLOSED and not obj.closed_at:
            obj.closed_at = timezone.now()
        super().save_model(request, obj, form, change)

    @admin.action(description='Назначить выбранные заявки на текущего пользователя')
    def assign_to_me(self, request, queryset):
        user = request.user
        updated = queryset.update(assignee=user)
        self.message_user(request, f'{updated} заявок(и) успешно назначены  на вас.')

    @admin.action(description="Массовое закрытие выбранных заявок")
    def close_tickets(self, request, queryset):
        now = timezone.now()
        updated = queryset.filter(status__in=[Ticket.Status.NEW, Ticket.Status.IN_PROGRESS, Ticket.Status.WAITING,]).update(status = Ticket.Status.CLOSED, closed_at=now)
        self.message_user(request, f'{updated} заявок(и) закрыто.')

    @admin.action(description='Выгрузка выбранных заявок в Excel')
    def export_selected(self, request, queryset):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Tickets'

        headers = ['ID', 'Тема', 'Клиент', 'Email', 'Описание', 'Приоритет', 'Статус', 'Исполнитель', 'Дата создания', 'Дата закрытия']
        sheet.append(headers)

        for col in 'ABCDEFGHIJ':
            sheet.column_dimensions[col].width = 18

        colors = {
            Ticket.Priority.LOW: '008000',
            Ticket.Priority.NORMAL: 'FFFF00',
            Ticket.Priority.HIGH: 'FFA500',
            Ticket.Priority.CRITICAL: 'FF0000'
        }

        for t in queryset:
            row = [
                t.external_id,
                t.subject,
                t.client.name or '',
                t.client.email or '',
                t.description or '',
                t.get_priority_display(),
                t.get_status_display(),
                str(t.assignee) or '',
                t.created_at.strftime('%Y-%m-%d %H:%M') if t.created_at else '',
                t.closed_at.strftime('%Y-%m-%d %H:%M') if t.closed_at else '',
            ]
            sheet.append(row)

            color = colors.get(t.priority)
            if color:
                cell = sheet.cell(row=sheet.max_row, column=6)
                cell.fill = PatternFill(start_color=color, end_color=color)
                cell.font = Font(bold=True, color='ffffff')
                cell.alignment = Alignment(horizontal='center')

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=tickets_export.xlsx'
        workbook.save(response)
        return response

    actions = [assign_to_me, close_tickets, export_selected]
    
    def get_urls(self):
        urls = super().get_urls()
        custom_url = [path('excel-import/', self.admin_site.admin_view(self.import_excel), name='ticket_import'), ]
        return custom_url + urls

    def import_excel(self, request):
        if request.method == 'POST':
            form = ExcelInputForm(request.POST, request.FILES)
            if form.is_valid():
                excel_file = request.FILES['excel_file']
                workbook = openpyxl.load_workbook(excel_file, data_only=True)
                sheet = workbook.active

                created_count = 0
                updated_count = 0
                skipped_count = 0
                errors = []
                total_rows = max(0, sheet.max_row - 1)

                for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                    if not row or not row[0]:
                        skipped_count += 1
                        errors.append(f'Строка {row_idx}: отсутствует ID задачи') 
                        continue
                    
                    external_id = str(row[0]).strip()
                    subject = str(row[1]) or ''
                    client_name = str(row[2]) or ''
                    client_email = str(row[3]) or ''
                    description = str(row[4]) or ''
                    priority = str(row[5]).strip().lower()
                    status = str(row[6]).strip().lower()
                    created_at = str(row[7])

                    if not subject:
                        skipped_count += 1
                        errors.append(f'Строка {row_idx}: отсутствует тема заявки')
                        continue

                    try:
                        priority_correct = Ticket.Priority(priority)
                    except ValueError:
                        skipped_count += 1
                        errors.append(f'Строка {row_idx}: неизвестный приоритет "{priority}"')
                        continue

                    try:
                        status_correct = Ticket.Status(status)
                    except ValueError:
                        skipped_count += 1
                        errors.append(f'Строка {row_idx}: неизвестный статус "{status}"')
                        continue

                    if created_at:
                        try:
                            created_at_correct = parse_datetime(str(created_at))
                        except Exception:
                            skipped_count += 1
                            errors.append(f'Строка {row_idx}: некорректная дата')
                            continue
                    else:
                        created_at_correct = None

                    client, _ = Client.objects.get_or_create(
                        email = client_email,
                        defaults={'name': client_name}
                    )

                    ticket, created = Ticket.objects.update_or_create(
                        external_id=external_id,
                        defaults={
                            'subject': subject,
                            'client': client,
                            'description': description,
                            'priority': priority_correct,
                            'status': status_correct,
                            'created_at': created_at_correct,
                        }
                    )
                    if created:
                        created_count += 1
                    else:
                        updated_count += 1

                ImportLog.objects.create(
                    user=request.user,
                    file_name=excel_file.name,
                    uploaded_at=timezone.now(),
                    total_rows=total_rows,
                    created_count=created_count,
                    updated_count=updated_count,
                    skipped_count=skipped_count,
                    errors=errors,
                )

                context = {
                    'created_count': created_count,
                    'updated_count': updated_count,
                    'skipped_count': skipped_count,
                    'errors': errors,
                }

                return render(request, 'helpdesk/admin/import_report.html', context)
        else:
            form = ExcelInputForm()
        
        context = {
            **self.admin_site.each_context(request),
            'form': form,
            'title': 'Импорт списка задач из Excel',
        }
        return render(request, 'helpdesk/admin/excel_import.html', context)